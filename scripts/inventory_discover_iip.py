#!/usr/bin/env python3
"""在庫発見器（経産省 鉱工業指数 IIP・業種／財別 原指数 月次・在庫／在庫率／出荷／生産）。

無料の一次統計だけから「数量側」の異変を列挙する。値段側の発見器（`driver_discover_boj.py`・日銀
CGPI/SPPI）と同じ型で、〈値上がり＋在庫減＝品薄／値下がり＋在庫増＝供給過剰〉を読む材料を
`output/inventory_discover.md` に出す（tasks/inventory_stat_entry.md・オーナー裁定 2026-09-21
「在庫の統計を入口に足す」）。

やること:
  1. 経産省の `b2020_zom1j.xlsx`（4シート: 生産／出荷／在庫／在庫率・業種・財別・原指数・月次）を取得し
     `data/iip/` に保存する（`--offline PATH` で保存済みファイルから読む）。
  2. 項目ごとに 在庫率・在庫・出荷 の〈前月比／3ヶ月比／前年同月比〉を出す。
  3. 在庫率の前年同月比で並べ、上位（在庫率↑＝供給過剰候補）と下位（在庫率↓＝品薄候補）を top N ずつ列挙。
     各行に出荷の前年比を並べ、〈在庫率↑かつ出荷↓〉〈在庫率↓かつ出荷↑〉を「整合」と印す。
  4. 項目名のトークンを center_pin（TOP1000）に部分一致させ候補会社を付ける（driver_discover_boj の
     match_companies を再利用・最大5社／項目・帰属 tier の判定はしない）。

やらないこと: 受益 tier の判定（beneficiary-attribution）・configs/*.json の編集・launchd 登録（別裁定）・
品目別ワークブック（`b2020_ygzosm1je.xlsx` 等・2026-09-21 は配信側が 202 空応答で未取得）。

fail-closed: 最新月が実行日から 4 ヶ月より古い／項目数が 50 未満／必須シートが無い → exit 2 で md を書かない。

使い方:
    python3 scripts/inventory_discover_iip.py [--top 15]
    python3 scripts/inventory_discover_iip.py --offline data/iip/b2020_zom1j.xlsx
    python3 scripts/inventory_discover_iip.py --selftest   # ネットワーク不要の固定データで検査
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import os
import re
import sys
import urllib.request
import zipfile
from typing import Any, Dict, List, Optional, Tuple

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

from driver_discover_boj import load_center_pin, match_companies  # noqa: E402
from monthly_sources import _ym_shift  # noqa: E402

IIP_URL = "https://www.meti.go.jp/statistics/tyo/iip/xls/b2020_zom1j.xlsx"
CACHE_DIR = os.path.join(ROOT, "data", "iip")
OUTPUT_PATH = os.path.join(ROOT, "output", "inventory_discover.md")
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_6) AppleWebKit/605.1.15 Safari/605.1.15"

SHEETS = ("在庫率", "在庫", "出荷", "生産")
REQUIRED_SHEETS = ("在庫率", "在庫", "出荷")
MIN_ITEMS = 50
MAX_STALE_MONTHS = 4
# 集計行（全業種の合計・財別の大枠）は「候補会社」を付けない＝社名に一致しても意味が無い
AGGREGATE_NAMES = {"鉱工業", "最終需要財", "投資財", "資本財", "資本財（除．輸送機械）", "建設財", "消費財",
                   "耐久消費財", "非耐久消費財", "生産財", "鉱工業用生産財", "その他用生産財", "製造工業"}


# ---------------------------------------------------------------- 取得・パース

def fetch_workbook(url: str = IIP_URL, cache_dir: str = CACHE_DIR) -> bytes:
    """xlsx を取得して cache_dir に保存し、bytes を返す（0 bytes は失敗として扱う）。"""
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = resp.read()
    if len(data) < 10_000 or data[:2] != b"PK":
        raise RuntimeError(f"取得失敗: {len(data)} bytes（xlsx でない・配信側の抑制の可能性）")
    os.makedirs(cache_dir, exist_ok=True)
    path = os.path.join(cache_dir, os.path.basename(url))
    with open(path, "wb") as f:
        f.write(data)
    return data


def parse_sheet(rows: List[Tuple[Any, ...]]) -> Tuple[List[str], List[Dict[str, Any]]]:
    """1シート（rows = 全行タプル）→ (月ラベル列, 項目リスト)。

    レイアウト（2026-09-21 実測）: 3行目 = ['品目番号','品目名称','ウエイト', '201801', ...]、
    4行目以降 = [code, name, weight, v1, v2, ...]。値は文字列か数値。空は None。
    """
    header_idx = None
    for i, r in enumerate(rows[:6]):
        if r and str(r[0] or "").strip() == "品目番号":
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("ヘッダ行（品目番号）が見つからない")
    header = rows[header_idx]
    # 年月にも速報の "P" が前置され得る（利用上の注意 A19「年月の左にPが表示されている場合は、速報値です」）
    months = [str(m).strip().lstrip("Pp").strip() for m in header[3:] if m is not None and str(m).strip()]
    items: List[Dict[str, Any]] = []
    for r in rows[header_idx + 1:]:
        if not r or r[0] is None or r[1] is None:
            continue
        code, name = str(r[0]).strip(), str(r[1]).strip()
        if not code or not name:
            continue
        series: Dict[str, float] = {}
        for m, v in zip(months, r[3:]):
            f = _num(v)
            if f is not None:
                series[m] = f
        items.append({"code": code, "name": name, "weight": _num(r[2]), "series": series})
    return months, items


def _num(v: Any) -> Optional[float]:
    if v is None:
        return None
    s = str(v).strip().lstrip("Pp").strip()  # 速報値は "P" 前置（利用上の注意）
    if s in ("", "-", "－", "…", "x", "X"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def load_all(data: bytes) -> Dict[str, Tuple[List[str], List[Dict[str, Any]]]]:
    import openpyxl  # 実行時 import（--selftest はネットワーク・openpyxl どちらも不要にする）

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: Dict[str, Tuple[List[str], List[Dict[str, Any]]]] = {}
    for sh in SHEETS:
        if sh in wb.sheetnames:
            out[sh] = parse_sheet(list(wb[sh].iter_rows(values_only=True)))
    missing = [s for s in REQUIRED_SHEETS if s not in out]
    if missing:
        raise ValueError(f"必須シートが無い: {missing}（sheetnames={wb.sheetnames}）")
    return out


# ---------------------------------------------------------------- 指標

def metrics(series: Dict[str, float], latest: str) -> Dict[str, Optional[float]]:
    """latest 月の値と 前月比／3ヶ月比／前年同月比（%）。欠測は None。"""
    cur = series.get(latest)

    def pct(back: int) -> Optional[float]:
        prev = series.get(_ym_shift(latest, -back))
        if cur is None or prev in (None, 0):
            return None
        return round((cur / prev - 1) * 100, 1)

    return {"value": cur, "mom": pct(1), "m3": pct(3), "yoy": pct(12)}


def latest_month(months: List[str], items: List[Dict[str, Any]]) -> str:
    """値が入っている最も新しい月（ヘッダ末尾が空欄のことがあるので実値で決める）。"""
    have = set()
    for it in items:
        have.update(it["series"].keys())
    cands = [m for m in months if m in have]
    if not cands:
        raise ValueError("値のある月が無い")
    return max(cands)


def parent_code(code: str) -> Optional[str]:
    """財別×業種の階層コード（例 3100010204）の親を返す。末尾から2桁ずつ 00 に戻す。"""
    for cut in (8, 6, 4):
        if code[cut:] != "0" * (len(code) - cut):
            return code[:cut] + "0" * (len(code) - cut)
    return None


def display_name(code: str, name: str, names: Dict[str, str]) -> str:
    """業種名は複数の財（資本財／建設財／消費財…）の下に同名で現れるので〈親の財／業種〉で表示する。"""
    par = parent_code(code)
    if par and par in names and code[8:] != "00":
        return f"{names[par]}／{name}"
    return name


def classify(r_yoy: Optional[float], s_yoy: Optional[float]) -> str:
    """二条件の判定（オーナー裁定の定義どおり）。片方だけの時は「候補」と言わず方向のみ書く。

    供給過剰候補 = 在庫率↑ かつ 出荷↓／品薄候補 = 在庫率↓ かつ 出荷↑。
    出荷が同方向・欠測なら「在庫率↑（出荷↑・要確認）」等に留める。
    """
    if r_yoy is None:
        return "判定不能"
    if r_yoy == 0:
        return "横ばい"
    if s_yoy is None:
        return ("在庫率↑" if r_yoy > 0 else "在庫率↓") + "（出荷 欠測・要確認）"
    if r_yoy > 0:
        return "供給過剰候補" if s_yoy < 0 else "在庫率↑（出荷↑・要確認）"
    return "品薄候補" if s_yoy > 0 else "在庫率↓（出荷↓・要確認）"


def build_table(sheets: Dict[str, Tuple[List[str], List[Dict[str, Any]]]]) -> Tuple[str, List[Dict[str, Any]]]:
    months, ratio_items = sheets["在庫率"]
    latest = latest_month(months, ratio_items)
    by_code = {s: {it["code"]: it for it in sheets[s][1]} for s in sheets}
    names = {it["code"]: it["name"] for it in ratio_items}
    rows: List[Dict[str, Any]] = []
    for it in ratio_items:
        code = it["code"]
        if not it["weight"]:
            continue  # ウエイト 0 = 指数を構成する品目が無い（利用上の注意）→ 除外
        row: Dict[str, Any] = {"code": code, "name": it["name"], "weight": it["weight"],
                               "display": display_name(code, it["name"], names)}
        for s in sheets:
            src = by_code[s].get(code)
            row[s] = metrics(src["series"], latest) if src else {"value": None, "mom": None, "m3": None, "yoy": None}
        row["label"] = classify(row["在庫率"]["yoy"], row["出荷"]["yoy"])
        row["aggregate"] = code[8:] == "00"  # 財・鉱工業の集計行（内訳と同じ枠で並べない）
        rows.append(row)
    return latest, rows


def rank(rows: List[Dict[str, Any]], top: int) -> Dict[str, List[Dict[str, Any]]]:
    scored = [r for r in rows if r["在庫率"]["yoy"] is not None and not r.get("aggregate")]
    up = sorted(scored, key=lambda r: -r["在庫率"]["yoy"])[:top]
    down = sorted(scored, key=lambda r: r["在庫率"]["yoy"])[:top]
    return {"up": [r for r in up if r["在庫率"]["yoy"] > 0], "down": [r for r in down if r["在庫率"]["yoy"] < 0]}


# ---------------------------------------------------------------- 出力

def _f(v: Optional[float]) -> str:
    return "—" if v is None else f"{v:+.1f}"


def _companies(row: Dict[str, Any], pins: List[Dict[str, Any]]) -> str:
    if row["name"] in AGGREGATE_NAMES or not pins:
        return "—"
    # 「輸送機械工業（除．自動車工業）」の除外語を検索語にしない（driver_discover_boj.item_tokens は "除" を落として
    # 肯定語にする＝Codex レビュー 2026-09-21 P2）。括弧内の「除．…」を先に切り落とす
    base = re.sub(r"[（(]除[．.].*?[）)]", "", row["name"])
    hits = match_companies(base, pins)
    return " / ".join(f"{h['code']} {h['name']}({h['sign']})" for h in hits) if hits else "—"


def render_markdown(run_date: str, latest: str, rows: List[Dict[str, Any]], ranked: Dict[str, List[Dict[str, Any]]],
                    pins: List[Dict[str, Any]], n_sheet_rows: int, source: str, top: int) -> str:
    n_labeled = sum(1 for r in rows if r["label"] != "判定不能")
    lines = [
        f"# 在庫発見器（IIP 業種・財別・原指数・月次）— {run_date} 実行",
        "",
        f"- データ源: {source}（2020年基準・在庫／在庫率／出荷／生産）",
        f"- **最新月: {latest}**（前年同月比は {_ym_shift(latest, -12)} 比・3ヶ月比は {_ym_shift(latest, -3)} 比）",
        f"- **母集団**: 在庫率シートの項目行 {n_sheet_rows} 行（除外= 番号・名称が空の行＋ウエイト 0 の行 {n_sheet_rows - len(rows)} 行）→ 取り込み {len(rows)} 項目・"
        f"ラベル付け {n_labeled}／{len(rows)} 項目（判定不能= 前年同月の欠測）",
        "- 読み方: **在庫率↑＋出荷↓＝供給過剰候補／在庫率↓＋出荷↑＝品薄候補**（二条件を満たす行だけ「候補」。片方だけは"
        "「要確認」）。数量側の候補であって、価格系列（日銀 CGPI・B2B 72 系列）との照合は未実施＝値上がり／値下がりの裏付けは別工程。"
        "候補会社は center_pin への語の部分一致（帰属 tier は未判定・beneficiary-attribution で裏取り）",
        "",
        "## 概況（集計行・財別）",
        "",
        "| 区分 | 在庫率 前年比 | 在庫 前年比 | 出荷 前年比 | 判定 |",
        "|---|---|---|---|---|",
        *[f"| {r['display']} | {_f(r['在庫率']['yoy'])} | {_f(r['在庫']['yoy'])} | {_f(r['出荷']['yoy'])} | {r['label']} |"
          for r in rows if r.get("aggregate")],
        "",
        f"## 在庫率が上がった内訳（業種×財・上位 {top}・集計行は除く）",
        "",
        "| 項目 | 在庫率 前年比 | 3ヶ月比 | 在庫 前年比 | 出荷 前年比 | 判定 | 候補会社 |",
        "|---|---|---|---|---|---|---|",
    ]
    for r in ranked["up"]:
        lines.append(f"| {r['display']} | {_f(r['在庫率']['yoy'])} | {_f(r['在庫率']['m3'])} | {_f(r['在庫']['yoy'])} | "
                     f"{_f(r['出荷']['yoy'])} | {r['label']} | {_companies(r, pins)} |")
    lines += ["", f"## 在庫率が下がった内訳（業種×財・上位 {top}・集計行は除く）", "",
              "| 項目 | 在庫率 前年比 | 3ヶ月比 | 在庫 前年比 | 出荷 前年比 | 判定 | 候補会社 |",
              "|---|---|---|---|---|---|---|"]
    for r in ranked["down"]:
        lines.append(f"| {r['display']} | {_f(r['在庫率']['yoy'])} | {_f(r['在庫率']['m3'])} | {_f(r['在庫']['yoy'])} | "
                     f"{_f(r['出荷']['yoy'])} | {r['label']} | {_companies(r, pins)} |")
    lines += ["", "## 注記", "",
              "- 原指数（季節調整なし）＝前年同月比で読む。前月比・3ヶ月比は季節の影響を含む。2020年基準の単一系列（201801〜）＝接続係数は未適用（基準改定をまたぐ比較はしない）",
              "- 最新月は速報値（P）を含むことがある＝翌月の確報で改定される",
              "- 業種・財別（約100項目）＝品目別（500超）は未取り込み（配信側の抑制で 2026-09-21 未取得・tasks/inventory_stat_entry.md）",
              "- launchd 未登録（手動実行のみ・定期化は別裁定）", ""]
    return "\n".join(lines)


# ---------------------------------------------------------------- 実行

def run(top: int, offline: Optional[str], out_path: str = OUTPUT_PATH) -> int:
    run_date = dt.date.today().isoformat()
    if offline:
        with open(offline, "rb") as f:
            data = f.read()
        source = offline
    else:
        data = fetch_workbook()
        source = IIP_URL
    sheets = load_all(data)
    n_sheet_rows = len(sheets["在庫率"][1])
    latest, rows = build_table(sheets)
    # fail-closed
    today = dt.date.today()
    age = (today.year * 12 + today.month) - (int(latest[:4]) * 12 + int(latest[4:6]))
    if age > MAX_STALE_MONTHS:
        print(f"[fail-closed] 最新月 {latest} が {age} ヶ月前＝stale（上限 {MAX_STALE_MONTHS}）", file=sys.stderr)
        return 2
    if len(rows) < MIN_ITEMS:
        print(f"[fail-closed] 項目数 {len(rows)} < {MIN_ITEMS}", file=sys.stderr)
        return 2
    n_ratio = sum(1 for r in rows if r["在庫率"]["yoy"] is not None)
    n_ship = sum(1 for r in rows if r["出荷"]["yoy"] is not None)
    if n_ratio < MIN_ITEMS * 0.8 or n_ship < MIN_ITEMS * 0.8:
        print(f"[fail-closed] 前年比の有効件数が不足: 在庫率 {n_ratio}／出荷 {n_ship}（下限 {int(MIN_ITEMS * 0.8)}）", file=sys.stderr)
        return 2
    ranked = rank(rows, top)
    pins = load_center_pin() if os.path.exists(os.path.join(ROOT, "data", "center_pin", "center_pin.jsonl")) else []
    md = render_markdown(run_date, latest, rows, ranked, pins, n_sheet_rows, source, top)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(md)
    print(f"wrote {out_path}")
    print(f"--- self-check: latest={latest} sheet_rows={n_sheet_rows} items={len(rows)} "
          f"labeled={sum(1 for r in rows if r['label'] != '判定不能')} up={len(ranked['up'])} down={len(ranked['down'])} "
          f"pins={len(pins)}")
    return 0


def _selftest() -> int:
    """固定データ: 在庫率↑＋出荷↓ の項目が供給過剰候補、在庫率↓＋出荷↑ が品薄候補になること。"""
    # 実データと同じ階層コード（財 3100010200 の下に業種 …0204 等）・同名の別財・ウエイト0・速報 "P" 付き年月
    months = ["202407", "202504", "202506", "P202507"]
    spec = [  # (code, name, weight)
        ("3000000000", "鉱工業", 100.0), ("3100010200", "資本財（除．輸送機械）", 50.0),
        ("3100010204", "電気・情報通信機械工業", 10.0), ("3100010201", "鉄鋼・非鉄金属工業", 10.0),
        ("3100010206", "輸送機械工業（除．自動車工業）", 10.0), ("3100020103", "電気・情報通信機械工業", 10.0),
        ("3100010303", "汎用・業務用機械工業", 0), ("3200000105", "電子部品・デバイス工業", 10.0),
    ]

    def sheet(vals: Dict[str, List[Any]]) -> List[Tuple[Any, ...]]:
        rows: List[Tuple[Any, ...]] = [("項目：テスト",), ("", "", "時系列コード"), ("品目番号", "品目名称", "ウエイト", *months)]
        for code, name, w in spec:
            rows.append((code, name, w, *vals[code]))
        return rows

    flat = [100, 100, 100, 100]
    ratio = sheet({"3000000000": flat, "3100010200": [100, 105, 108, 110], "3100010204": [100, 110, 120, "P130"],
                   "3100010201": [100, 95, 90, 80], "3100010206": [100, 100, 100, 105], "3100020103": [100, 100, 100, 90],
                   "3100010303": flat, "3200000105": [None, 90, 90, 90]})
    stock = sheet({c: flat for c, _, _ in spec})
    ship = sheet({"3000000000": flat, "3100010200": [100, 99, 98, 97], "3100010204": [100, 96, 92, 90],
                  "3100010201": [100, 105, 110, 112], "3100010206": [100, 100, 100, 103], "3100020103": [100, 100, 100, 95],
                  "3100010303": flat, "3200000105": flat})
    sheets = {"在庫率": parse_sheet(ratio), "在庫": parse_sheet(stock), "出荷": parse_sheet(ship)}
    assert sheets["在庫率"][0][-1] == "202507", sheets["在庫率"][0]  # 年月の P を落とす
    latest, rows = build_table(sheets)
    assert latest == "202507", latest
    assert len(rows) == 7, [r["code"] for r in rows]  # ウエイト 0 の 3100010303 を除外
    by = {r["code"]: r for r in rows}
    assert by["3100010204"]["display"] == "資本財（除．輸送機械）／電気・情報通信機械工業", by["3100010204"]["display"]
    assert by["3100010204"]["在庫率"]["yoy"] == 30.0 and by["3100010204"]["label"] == "供給過剰候補"
    assert by["3100010201"]["label"] == "品薄候補", by["3100010201"]["label"]
    assert by["3100010206"]["label"] == "在庫率↑（出荷↑・要確認）", by["3100010206"]["label"]   # 片方だけ＝候補と言わない
    assert by["3100020103"]["label"] == "在庫率↓（出荷↓・要確認）", by["3100020103"]["label"]
    assert by["3200000105"]["label"] == "判定不能" and by["3000000000"]["label"] == "横ばい"
    assert by["3100010200"]["aggregate"] and by["3100010200"]["label"] == "供給過剰候補"
    ranked = rank(rows, 5)
    assert [r["code"] for r in ranked["up"]] == ["3100010204", "3100010206"], ranked["up"]   # 集計行 3100010200 は入らない
    assert [r["code"] for r in ranked["down"]] == ["3100010201", "3100020103"], ranked["down"]
    pins = [{"code": "7211", "name": "三菱自動車工業", "pin": "自動車", "note": "", "sign": "+"}]
    assert _companies(by["3100010206"], pins) == "—", _companies(by["3100010206"], pins)  # 「除．自動車工業」を検索語にしない
    md = render_markdown("2026-09-21", latest, rows, ranked, pins, 8, "selftest", 5)
    assert "母集団" in md and "## 概況（集計行・財別）" in md and "資本財（除．輸送機械） | +10.0" in md
    assert "3ヶ月比は 202504 比" in md
    print("selftest OK: latest=202507 sheet_rows=8 items=7 labeled=6 up=2 down=2 aggregate=2")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--top", type=int, default=15)
    ap.add_argument("--offline", help="保存済み xlsx から読む（ネットワーク不要）")
    ap.add_argument("--out", default=OUTPUT_PATH)
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return _selftest()
    try:
        return run(a.top, a.offline, a.out)
    except (ValueError, RuntimeError, OSError, zipfile.BadZipFile) as e:  # 取得失敗・シート名変更・0 bytes = 品質検査の失敗として exit 2
        print(f"[fail-closed] {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
